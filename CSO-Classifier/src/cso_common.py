from __future__ import annotations

import json
import math
from pathlib import Path

from cso_classifier import CSOClassifier
from cso_classifier.ontology import Ontology
from openpyxl import load_workbook

SCORE = 100
TOP_K = 7

PROPOSAL_INTENDED_END_USER_ROLE = "Learner"
PROPOSAL_CONTEXT = "Higher Education"
PROPOSAL_TYPICAL_AGE_RANGE = "18+"

INTERACTIVITY_TYPE_ALLOWED_LEVELS = {
    "Expositive": {"Very Low"},
    "Mixed": {"Low"},
    "Active": {"Medium", "High", "Very High"},
}

IMPOSSIBLE_SEMANTIC_DENSITY_DIFFICULTY = {
    ("Very Low", "Difficult"),
    ("Very Low", "Very Difficult"),
    ("Low", "Very Difficult"),
    ("High", "Very Easy"),
    ("Very High", "Very Easy"),
    ("Very High", "Easy"),
}

HEADER_ALIASES = {
    "interactivityType": ["interactivityType", "Interactivity_Type", "Interactivity Type"],
    "learningResourceType": ["learningResourceType", "Learning_Resource_Type", "Learning Resource Type"],
    "interactivityLevel": ["interactivityLevel", "Interactivity_Level", "Interactivity Level"],
    "semanticDensity": ["semanticDensity", "Semantic_Density", "Semantic Density"],
    "intendedEndUserRole": [
        "intendedEndUserRole",
        "IntendedEndUserRole",
        "Intended_End_User_Role",
        "Intended End User Role",
    ],
    "context": ["context", "Context"],
    "typicalAgeRange": [
        "typicalAgeRange",
        "TypicalAgeRange",
        "Typical_Age_Range",
        "Typical Age Range",
    ],
    "difficulty": ["difficulty", "Difficulty"],
    "typicalLearningTime": [
        "typicalLearningTime",
        "Typical_Learning_Time ( minutes",
        "Typical_Learning_Time (minutes)",
        "Typical_Learning_Time",
        "Typical Learning Time (minutes)",
        "Typical Learning Time",
    ],
    "description": ["description", "Description"],
    "language": ["language", "Language"],
}


def uniq_sorted(values):
    return sorted(set(values))


def self_arrow(values):
    return "\n".join(f"{value}->{value}({SCORE})" for value in values)


def normalize_header(value):
    if value is None:
        return ""
    return str(value).strip().lower()


def find_column_by_header(ws, possible_names):
    targets = {normalize_header(name) for name in possible_names}
    for cell in ws[1]:
        if normalize_header(cell.value) in targets:
            return cell.column
    return None


def last_nonempty_header_column(ws):
    last_col = 0
    for cell in ws[1]:
        if cell.value is not None and str(cell.value).strip():
            last_col = cell.column
    return last_col


def get_or_create_column(ws, header_name):
    col = find_column_by_header(ws, [header_name])
    if col is not None:
        return col

    new_col = last_nonempty_header_column(ws) + 1
    ws.cell(row=1, column=new_col).value = header_name
    return new_col


def delete_columns_by_headers(ws, header_names):
    targets = {normalize_header(name) for name in header_names}

    while True:
        col_to_delete = None
        for cell in ws[1]:
            if normalize_header(cell.value) in targets:
                col_to_delete = cell.column
                break

        if col_to_delete is None:
            break

        ws.delete_cols(col_to_delete)


def normalize_genolom_headers(ws):
    for canonical_name, aliases in HEADER_ALIASES.items():
        col = find_column_by_header(ws, aliases)
        if col is not None:
            ws.cell(row=1, column=col).value = canonical_name


def ensure_column_after(ws, header_name, after_header):
    existing = find_column_by_header(ws, [header_name])
    if existing is not None:
        return existing

    after_col = find_column_by_header(ws, [after_header])
    if after_col is None:
        raise ValueError(
            f"Cannot insert {header_name!r} because {after_header!r} was not found."
        )

    new_col = after_col + 1
    ws.insert_cols(new_col)
    ws.cell(row=1, column=new_col).value = header_name
    return new_col


def set_column_value_for_all_rows(ws, header_name, value):
    col = find_column_by_header(ws, [header_name])
    if col is None:
        raise ValueError(f"Required column {header_name!r} was not found.")

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=col).value = value


def normalize_interactivity_level(interactivity_type, current_level):
    """Return the closest proposal-valid interactivityLevel.

    Proposal Table 3:
      Expositive -> Very Low
      Mixed      -> Low
      Active     -> Medium / High / Very High

    For an invalid Active level, Medium is the minimum valid level and is
    therefore the minimum-change repair.
    """
    interactivity_type = str(interactivity_type).strip()
    current_level = str(current_level).strip()

    if interactivity_type == "Expositive":
        return "Very Low"
    if interactivity_type == "Mixed":
        return "Low"
    if interactivity_type == "Active":
        if current_level in INTERACTIVITY_TYPE_ALLOWED_LEVELS["Active"]:
            return current_level
        return "Medium"

    return current_level


def normalize_proposal_metadata(ws):
    normalize_genolom_headers(ws)

    ensure_column_after(ws, "intendedEndUserRole", "semanticDensity")
    ensure_column_after(ws, "context", "intendedEndUserRole")
    ensure_column_after(ws, "typicalAgeRange", "context")

    set_column_value_for_all_rows(
        ws, "intendedEndUserRole", PROPOSAL_INTENDED_END_USER_ROLE
    )
    set_column_value_for_all_rows(ws, "context", PROPOSAL_CONTEXT)
    set_column_value_for_all_rows(
        ws, "typicalAgeRange", PROPOSAL_TYPICAL_AGE_RANGE
    )

    type_col = find_column_by_header(ws, ["interactivityType"])
    level_col = find_column_by_header(ws, ["interactivityLevel"])
    if type_col is None or level_col is None:
        raise ValueError("interactivityType/interactivityLevel columns were not found.")

    changed_rows = 0
    for row in range(2, ws.max_row + 1):
        current_type = ws.cell(row=row, column=type_col).value
        current_level = ws.cell(row=row, column=level_col).value
        normalized_level = normalize_interactivity_level(current_type, current_level)
        if normalized_level != current_level:
            ws.cell(row=row, column=level_col).value = normalized_level
            changed_rows += 1

    return changed_rows


def prepare_genolom_columns(ws):
    changed_rows = normalize_proposal_metadata(ws)

    delete_columns_by_headers(
        ws,
        [
            "simple_matches",
            "extended_matches",
            "matched_phrases",
            "Final_Keywords",
            "keywords",
            "Keyword_Depths",
            "Depth",
        ],
    )

    keywords_col = ensure_column_after(ws, "keywords", "language")
    depth_col = ensure_column_after(ws, "Depth", "keywords")
    return keywords_col, depth_col, changed_rows


def get_topic_depth(topic, cso, depth_cache):
    if topic is None:
        return "NA"

    topic = str(topic).strip()
    if not topic:
        return "NA"

    if topic in depth_cache:
        return depth_cache[topic]

    root_topic = "computer science"
    graph = cso.get_ontology_graph()

    candidate_topics = []
    try:
        candidate_topics.append(cso.get_primary_label(topic))
    except Exception:
        pass

    candidate_topics.append(topic)
    candidate_topics = list(dict.fromkeys(candidate_topics))

    best_depth = None
    for candidate in candidate_topics:
        try:
            dist = graph.shortest_paths_dijkstra(root_topic, candidate)[0][0]
        except Exception:
            try:
                dist = graph.distances(
                    source=root_topic,
                    target=candidate,
                    mode="ALL",
                )[0][0]
            except Exception:
                continue

        if dist is not None and dist != float("inf"):
            try:
                if not math.isinf(dist):
                    dist = int(dist)
                    if best_depth is None or dist < best_depth:
                        best_depth = dist
            except Exception:
                pass

    depth_cache[topic] = "NA" if best_depth is None else best_depth
    return depth_cache[topic]


def keyword_depths(topics, cso, depth_cache):
    lines = []
    for topic in topics:
        depth = get_topic_depth(topic, cso, depth_cache)
        lines.append(f"{topic}>{depth}")
    return "\n".join(lines)


def extract_outputs(text, clf, cso, depth_cache, include_extended):
    if text is None or not str(text).strip():
        matched = {"syntactic": [], "union": []}
        if include_extended:
            matched = {
                "syntactic": [],
                "semantic": [],
                "enhanced": [],
                "union": [],
            }
        return {
            "simple_matches": "",
            "extended_matches": "",
            "matched_phrases": json.dumps(matched, ensure_ascii=False),
            "keywords": "",
            "depth": "",
        }

    result = clf.run(str(text))
    syntactic = uniq_sorted(result.get("syntactic", []))

    if include_extended:
        semantic = uniq_sorted(result.get("semantic", []))
        enhanced = uniq_sorted(result.get("enhanced", []))
        union = uniq_sorted(result.get("union", []))

        syntactic_set = set(syntactic)
        extended = uniq_sorted(
            topic
            for topic in semantic + enhanced
            if topic not in syntactic_set
        )

        top_matched = syntactic[:TOP_K]
        if len(top_matched) < TOP_K:
            top_matched += extended[: TOP_K - len(top_matched)]

        matched_phrases = json.dumps(
            {
                "syntactic": syntactic,
                "semantic": semantic,
                "enhanced": enhanced,
                "union": union,
            },
            ensure_ascii=False,
        )
    else:
        extended = []
        top_matched = syntactic[:TOP_K]
        matched_phrases = json.dumps(
            {"syntactic": syntactic, "union": syntactic},
            ensure_ascii=False,
        )

    return {
        "simple_matches": self_arrow(syntactic),
        "extended_matches": self_arrow(extended),
        "matched_phrases": matched_phrases,
        "keywords": "\n".join(top_matched),
        "depth": keyword_depths(top_matched, cso, depth_cache),
    }


def validate_proposal_constraints(ws):
    required = [
        "interactivityType",
        "interactivityLevel",
        "semanticDensity",
        "difficulty",
        "intendedEndUserRole",
        "context",
        "typicalAgeRange",
    ]
    columns = {name: find_column_by_header(ws, [name]) for name in required}
    missing = [name for name, col in columns.items() if col is None]
    if missing:
        raise ValueError("Missing proposal columns: " + ", ".join(missing))

    errors = []
    for row in range(2, ws.max_row + 1):
        interactivity_type = ws.cell(row=row, column=columns["interactivityType"]).value
        interactivity_level = ws.cell(row=row, column=columns["interactivityLevel"]).value
        semantic_density = ws.cell(row=row, column=columns["semanticDensity"]).value
        difficulty = ws.cell(row=row, column=columns["difficulty"]).value

        allowed = INTERACTIVITY_TYPE_ALLOWED_LEVELS.get(interactivity_type)
        if allowed is not None and interactivity_level not in allowed:
            errors.append(
                f"Row {row}: invalid interactivity pair "
                f"({interactivity_type!r}, {interactivity_level!r})"
            )

        if (semantic_density, difficulty) in IMPOSSIBLE_SEMANTIC_DENSITY_DIFFICULTY:
            errors.append(
                f"Row {row}: impossible semanticDensity/difficulty pair "
                f"({semantic_density!r}, {difficulty!r})"
            )

        fixed = {
            "intendedEndUserRole": PROPOSAL_INTENDED_END_USER_ROLE,
            "context": PROPOSAL_CONTEXT,
            "typicalAgeRange": PROPOSAL_TYPICAL_AGE_RANGE,
        }
        for name, expected in fixed.items():
            actual = ws.cell(row=row, column=columns[name]).value
            if actual != expected:
                errors.append(
                    f"Row {row}: {name}={actual!r}; expected {expected!r}"
                )

    if errors:
        preview = "\n".join(errors[:20])
        suffix = "" if len(errors) <= 20 else f"\n... and {len(errors) - 20} more"
        raise ValueError("Proposal validation failed:\n" + preview + suffix)


def run_cso_pipeline(input_file: Path, output_file: Path, *, include_extended: bool):
    wb = load_workbook(input_file)
    ws = wb.active

    keywords_col, depth_col, normalized_rows = prepare_genolom_columns(ws)

    if not include_extended:
        delete_columns_by_headers(ws, ["extended_matches"])

    text_col = find_column_by_header(
        ws,
        ["Raw Text", "text", "paragraph", "content", "Paragraph Text"],
    )
    if text_col is None:
        raise ValueError("Raw Text column was not found.")

    simple_col = get_or_create_column(ws, "simple_matches")
    extended_col = (
        get_or_create_column(ws, "extended_matches") if include_extended else None
    )
    matched_col = get_or_create_column(ws, "matched_phrases")

    clf = CSOClassifier(use_full_model=False)
    cso = Ontology(load_ontology=True, silent=True)
    depth_cache = {}

    total_rows = ws.max_row - 1
    for row in range(2, ws.max_row + 1):
        text = ws.cell(row=row, column=text_col).value
        out = extract_outputs(
            text,
            clf,
            cso,
            depth_cache,
            include_extended=include_extended,
        )

        ws.cell(row=row, column=keywords_col).value = out["keywords"]
        ws.cell(row=row, column=depth_col).value = out["depth"]
        ws.cell(row=row, column=simple_col).value = out["simple_matches"]
        if extended_col is not None:
            ws.cell(row=row, column=extended_col).value = out["extended_matches"]
        ws.cell(row=row, column=matched_col).value = out["matched_phrases"]

        print(f"Processed row {row - 1}/{total_rows}")

    validate_proposal_constraints(ws)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    print("\nDone.")
    print(f"Saved: {output_file}")
    print(f"Proposal-normalized interactivity rows: {normalized_rows}")
    print(
        "Fixed chapter metadata: "
        f"intendedEndUserRole={PROPOSAL_INTENDED_END_USER_ROLE}, "
        f"context={PROPOSAL_CONTEXT}, "
        f"typicalAgeRange={PROPOSAL_TYPICAL_AGE_RANGE}"
    )
