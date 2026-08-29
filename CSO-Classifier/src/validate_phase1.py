from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

FILES = {
    "annotated": PROJECT_ROOT / "data" / "ch01_annotated.xlsx",
    "main": PROJECT_ROOT / "data" / "ch01_cso_main.xlsx",
    "extended": PROJECT_ROOT / "data" / "ch01_cso_extended.xlsx",
}

FIXED = {
    "intendedEndUserRole": "Learner",
    "context": "Higher Education",
    "typicalAgeRange": "18+",
}

ALLOWED_LEVELS = {
    "Expositive": {"Very Low"},
    "Mixed": {"Low"},
    "Active": {"Medium", "High", "Very High"},
}

IMPOSSIBLE_SD_DIFFICULTY = {
    ("Very Low", "Difficult"),
    ("Very Low", "Very Difficult"),
    ("Low", "Very Difficult"),
    ("High", "Very Easy"),
    ("Very High", "Very Easy"),
    ("Very High", "Easy"),
}


def validate(path: Path, kind: str):
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active

    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }

    required = {
        "Paragraph #",
        "interactivityType",
        "interactivityLevel",
        "semanticDensity",
        "intendedEndUserRole",
        "context",
        "typicalAgeRange",
        "difficulty",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise AssertionError(f"{path.name}: missing columns {missing}")

    data_rows = 0
    keyword_rows = 0
    errors = []

    for row in range(2, ws.max_row + 1):
        paragraph_id = ws.cell(row, headers["Paragraph #"]).value
        if paragraph_id is None:
            continue
        data_rows += 1

        t = ws.cell(row, headers["interactivityType"]).value
        level = ws.cell(row, headers["interactivityLevel"]).value
        allowed = ALLOWED_LEVELS.get(t)
        if allowed is not None and level not in allowed:
            errors.append(
                f"row {row}: invalid interactivity pair ({t}, {level})"
            )

        sd = ws.cell(row, headers["semanticDensity"]).value
        diff = ws.cell(row, headers["difficulty"]).value
        if (sd, diff) in IMPOSSIBLE_SD_DIFFICULTY:
            errors.append(
                f"row {row}: impossible semanticDensity/difficulty ({sd}, {diff})"
            )

        for column, expected in FIXED.items():
            actual = ws.cell(row, headers[column]).value
            if actual != expected:
                errors.append(
                    f"row {row}: {column}={actual!r}, expected {expected!r}"
                )

        if "keywords" in headers:
            value = ws.cell(row, headers["keywords"]).value
            if value is not None and str(value).strip():
                keyword_rows += 1

    if data_rows != 298:
        errors.append(f"expected 298 rows, found {data_rows}")

    if kind == "main" and "extended_matches" in headers:
        errors.append("main output must not contain extended_matches")

    if kind == "extended" and "extended_matches" not in headers:
        errors.append("extended output must contain extended_matches")

    wb.close()

    if errors:
        raise AssertionError(
            f"{path.name}: validation failed\n" + "\n".join(errors[:30])
        )

    print(
        f"PASS {path.name}: rows={data_rows}, "
        f"keyword_rows={keyword_rows if 'keywords' in headers else 'n/a'}"
    )


if __name__ == "__main__":
    for kind, path in FILES.items():
        validate(path, kind)
