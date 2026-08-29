from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook
from ..models.genome import GenoLOMGenome
from ..models.individual import Individual

REQUIRED_HEADERS = {
    "Paragraph #", "Type", "Raw Text", "Page #",
    "interactivityType", "learningResourceType", "interactivityLevel", "semanticDensity",
    "intendedEndUserRole", "context", "typicalAgeRange", "difficulty",
    "typicalLearningTime", "description", "language", "keywords",
}


def load_generation0(path: str | Path) -> list[Individual]:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
    except StopIteration:
        raise ValueError("Generation 0 workbook is empty.")

    missing = REQUIRED_HEADERS.difference(headers)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")

    out: list[Individual] = []
    for excel_index, values in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() != "" for v in values):
            continue
        row = dict(zip(headers, values))
        source_id = str(row.get("Paragraph #") or len(out) + 1).strip()
        genome = GenoLOMGenome.from_mapping(row)
        out.append(Individual(
            individual_id=f"G0_LO_{len(out)+1:03d}",
            content=str(row.get("Raw Text") or ""),
            genome=genome,
            source_type=str(row.get("Type") or ""),
            source_page=str(row.get("Page #") or ""),
            source_row_id=source_id,
            generation_created=0,
            created_by="initial_population",
            extra={
                "excel_row": excel_index,
                "Depth": row.get("Depth"),
                "simple_matches": row.get("simple_matches"),
                "matched_phrases": row.get("matched_phrases"),
            },
        ))
    return out
